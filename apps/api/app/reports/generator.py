"""Report generation: Jinja2 + WeasyPrint for PDF, openpyxl for XLSX, plain
JSON for the third format. No client-side rendering (html2canvas/jspdf) —
per PROMPT_CLAUDE_CODE.md §3, reports are produced server-side from the
same engine output the UI shows.
"""

from __future__ import annotations

import json
from io import BytesIO
from pathlib import Path
from typing import Any

from jinja2 import Environment, FileSystemLoader
from openpyxl import Workbook
from weasyprint import HTML  # type: ignore[import-untyped]  # no py.typed marker upstream

from ecobalance_core import reference

_TEMPLATES_DIR = Path(__file__).resolve().parent / "templates"
_env = Environment(loader=FileSystemLoader(str(_TEMPLATES_DIR)), autoescape=True)

_RATING_LABELS = {"critical": "Критический", "significant": "Значимый", "low": "Низкий"}


def _has_demo_factors(result: dict[str, Any]) -> bool:
    return any(flag["code"] == "demo_factors" for flag in result.get("data_quality_flags", []))


def render_pcf_pdf(
    *, company_name: str, product_name: str, result: dict[str, Any], generated_at: str
) -> bytes:
    template = _env.get_template("pcf_report.html.jinja2")
    html = template.render(
        company_name=company_name,
        product_name=product_name,
        boundary=result["boundary"],
        functional_unit=result["functional_unit"],
        gwp_set=result["gwp_set"],
        pcf_full=result["pcf_full_kg_co2_per_unit"],
        scope1=result["scope1_kg_co2_per_unit"],
        scope2=result["scope2_kg_co2_per_unit"],
        scope3=result["scope3_kg_co2_per_unit"],
        hotspots=result["hotspots"],
        rating_labels=_RATING_LABELS,
        data_quality_score=result["data_quality_score"],
        data_quality_flags=result["data_quality_flags"],
        allocation_factor=result["allocation_factor"],
        has_demo_factors=_has_demo_factors(result),
        generated_at=generated_at,
        factor_version="demo-2024",
    )
    pdf_bytes: bytes = HTML(string=html).write_pdf()
    return pdf_bytes


def render_cbam_pdf(
    *,
    company_name: str,
    product_name: str,
    product_id: str,
    route_name: str,
    period: str,
    result: dict[str, Any],
    ets_price_eur_per_t: float,
    default_markup_pct: float,
    generated_at: str,
) -> bytes:
    product = reference.get_product(product_id)
    embedded_kg = float(result["cbam_embedded_kg_co2_per_unit"] or 0)
    embedded_t = embedded_kg / 1000
    default_embedded_t = embedded_t * (1 + default_markup_pct / 100)

    template = _env.get_template("cbam_report.html.jinja2")
    html = template.render(
        company_name=company_name,
        product_name=product_name,
        cn_code=product.cn_code or "—",
        route_name=route_name,
        period=period,
        functional_unit=result["functional_unit"],
        cbam_embedded=embedded_kg,
        cbam_embedded_t=round(embedded_t, 3),
        default_embedded_t=round(default_embedded_t, 3),
        scope1=result["scope1_kg_co2_per_unit"],
        scope2=result["scope2_kg_co2_per_unit"],
        upstream=next(
            (
                b["value_kg_co2_per_unit"]
                for b in result["breakdown"]
                if b["category"] == "upstream"
            ),
            0,
        ),
        ets_price=ets_price_eur_per_t,
        cost_actual=round(embedded_t * ets_price_eur_per_t, 2),
        cost_default=round(default_embedded_t * ets_price_eur_per_t, 2),
        data_quality_score=result["data_quality_score"],
        data_quality_flags=result["data_quality_flags"],
        has_demo_factors=_has_demo_factors(result),
        generated_at=generated_at,
        factor_version="demo-2024",
    )
    pdf_bytes: bytes = HTML(string=html).write_pdf()
    return pdf_bytes


def render_scenario_pdf(
    *,
    company_name: str,
    product_name: str,
    scenario_name: str,
    baseline: dict[str, Any],
    scenario: dict[str, Any],
    measures_meta: list[dict[str, Any]],
    generated_at: str,
) -> bytes:
    template = _env.get_template("scenario_report.html.jinja2")
    cbam_applicable = baseline.get("cbam_applicable", False)
    delta_pcf = float(baseline["pcf_full_kg_co2_per_unit"]) - float(
        scenario["pcf_full_kg_co2_per_unit"]
    )
    delta_pct = (
        (delta_pcf / float(baseline["pcf_full_kg_co2_per_unit"]) * 100)
        if float(baseline["pcf_full_kg_co2_per_unit"]) != 0
        else 0
    )
    html = template.render(
        company_name=company_name,
        product_name=product_name,
        scenario_name=scenario_name,
        baseline_pcf=baseline["pcf_full_kg_co2_per_unit"],
        scenario_pcf=scenario["pcf_full_kg_co2_per_unit"],
        delta_pcf=round(delta_pcf, 2),
        delta_pct=round(delta_pct, 1),
        cbam_applicable=cbam_applicable,
        baseline_cbam=baseline.get("cbam_embedded_kg_co2_per_unit"),
        scenario_cbam=scenario.get("cbam_embedded_kg_co2_per_unit"),
        delta_cbam=(
            round(
                float(baseline["cbam_embedded_kg_co2_per_unit"])
                - float(scenario["cbam_embedded_kg_co2_per_unit"]),
                2,
            )
            if cbam_applicable
            else None
        ),
        measures=measures_meta,
        has_demo_factors=_has_demo_factors(scenario),
        generated_at=generated_at,
        factor_version="demo-2024",
    )
    pdf_bytes: bytes = HTML(string=html).write_pdf()
    return pdf_bytes


def render_result_xlsx(*, company_name: str, product_name: str, result: dict[str, Any]) -> bytes:
    workbook = Workbook()
    summary = workbook.active
    assert summary is not None  # a freshly created Workbook always has an active sheet
    summary.title = "Результат"
    summary.append(["Компания", company_name])
    summary.append(["Продукт", product_name])
    summary.append(["PCF (кг CO2-экв./ед.)", float(result["pcf_full_kg_co2_per_unit"])])
    if result.get("cbam_embedded_kg_co2_per_unit") is not None:
        summary.append(
            ["CBAM embedded (кг CO2-экв./ед.)", float(result["cbam_embedded_kg_co2_per_unit"])]
        )
    summary.append(["Scope 1", float(result["scope1_kg_co2_per_unit"])])
    summary.append(["Scope 2", float(result["scope2_kg_co2_per_unit"])])
    summary.append(["Scope 3", float(result["scope3_kg_co2_per_unit"])])

    breakdown_sheet = workbook.create_sheet("Разбивка")
    breakdown_sheet.append(["Категория", "Scope", "Значение, кг CO2-экв.", "Доля, %"])
    for item in result["breakdown"]:
        breakdown_sheet.append(
            [
                item["label"],
                item["scope"],
                float(item["value_kg_co2_per_unit"]),
                float(item["percentage"]),
            ]
        )

    hotspots_sheet = workbook.create_sheet("Hotspots")
    hotspots_sheet.append(["Источник", "Scope", "Вклад, %", "Оценка"])
    for hotspot in result["hotspots"]:
        hotspots_sheet.append(
            [
                hotspot["label"],
                hotspot["scope"],
                float(hotspot["contribution_pct"]),
                hotspot["rating"],
            ]
        )

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def render_result_json(*, company_name: str, product_name: str, result: dict[str, Any]) -> bytes:
    payload = {"company_name": company_name, "product_name": product_name, "result": result}
    return json.dumps(payload, ensure_ascii=False, indent=2).encode("utf-8")
